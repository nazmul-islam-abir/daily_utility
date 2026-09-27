from pptx import Presentation

p = Presentation('app_pptx/Orange Modern Nutrition & Meal App Pitch Deck Presentation.pptx')
for i, slide in enumerate(p.slides, 1):
    print(f'\n===== Slide {i} =====')
    for shape in slide.shapes:
        try:
            l = round(shape.left / 914400, 2)
            t = round(shape.top / 914400, 2)
            w = round(shape.width / 914400, 2)
            h = round(shape.height / 914400, 2)
        except Exception:
            l = t = w = h = 0
        txt = shape.text_frame.text[:60] if shape.has_text_frame else ''
        print(f'  {shape.name:18s} L={l:5.2f} T={t:5.2f} W={w:5.2f} H={h:5.2f} | {txt}')
