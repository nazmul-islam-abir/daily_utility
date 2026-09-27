from pptx import Presentation
from pptx.dml.color import RGBColor

p = Presentation('app_pptx/Orange Modern Nutrition & Meal App Pitch Deck Presentation.pptx')

# Sample colors from the first slide across freeform / group / text shapes
def sample(shape, depth=0):
    indent = '  ' * depth
    name = shape.name
    extra = ''
    try:
        if shape.fill.type == 1:  # SOLID
            try:
                extra = f" fill={shape.fill.fore_color.rgb}"
            except Exception:
                pass
    except Exception:
        pass
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                try:
                    extra += f" text='{run.text[:30]}' c={run.font.color.rgb}"
                except Exception:
                    pass
                break
            break
    print(f"{indent}{name}{extra}")
    if shape.shape_type == 6:  # GROUP
        for sh in shape.shapes:
            sample(sh, depth + 1)

for sh in p.slides[0].shapes:
    sample(sh)
