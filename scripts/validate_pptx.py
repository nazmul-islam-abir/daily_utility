"""Validate the rebuilt Daily Utility PPTX: dump all text frames per slide."""
import sys
from pptx import Presentation

OUT = 'app_pptx/Daily Utility App Pitch Deck.pptx'
prs = Presentation(OUT)

# Force UTF-8 output so we can see em-dashes, বাংলা text, etc.
sys.stdout.reconfigure(encoding='utf-8')

print(f'File: {OUT}')
print(f'Slide size: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in')
print(f'Slide count: {len(prs.slides)}\n')

for i, slide in enumerate(prs.slides, 1):
    print(f'\n===== Slide {i} =====')
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            txt = shape.text_frame.text.replace('\n', ' / ')
            print(f'  [{shape.name}] {txt}')
        if shape.shape_type == 13:  # PICTURE
            print(f'  [PICTURE {shape.name}] w={shape.width/914400:.2f}in '
                  f'h={shape.height/914400:.2f}in')

print('\nValidation complete.')
